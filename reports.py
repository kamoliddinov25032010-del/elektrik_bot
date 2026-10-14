from datetime import datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

import db
from utils import STATUS_LABEL, fmt


def cur_month() -> str:
    return datetime.now().strftime("%Y-%m")


def prev_month() -> str:
    d = datetime.now().replace(day=1) - timedelta(days=1)
    return d.strftime("%Y-%m")


def _u(col, usta):
    return (f" AND {col}=?", [usta]) if usta else ("", [])


async def summary(month: str, usta=None) -> dict:
    su, pu = _u("o.usta_id", usta)
    income = await db.scalar(
        "SELECT COALESCE(SUM(p.amount),0) FROM payments p JOIN orders o ON o.id=p.order_id "
        "WHERE strftime('%Y-%m', p.created_at)=?" + su, [month] + pu)
    mats = await db.scalar(
        "SELECT COALESCE(SUM(m.qty*m.price),0) FROM materials m JOIN orders o ON o.id=m.order_id "
        "WHERE strftime('%Y-%m', m.created_at)=?" + su, [month] + pu)
    se, pe = _u("by_user", usta)
    exp = await db.scalar(
        "SELECT COALESCE(SUM(amount),0) FROM expenses WHERE strftime('%Y-%m', created_at)=?" + se, [month] + pe)
    done = await db.scalar(
        "SELECT COUNT(*) FROM orders o WHERE status IN ('tayyor','topshirildi') "
        "AND strftime('%Y-%m', updated_at)=?" + su, [month] + pu)
    new = await db.scalar(
        "SELECT COUNT(*) FROM orders o WHERE strftime('%Y-%m', created_at)=?" + su, [month] + pu)
    debt = await db.scalar(
        f"SELECT COALESCE(SUM(wp+mat_sum-paid),0) FROM ({db.ORDER_SELECT}) o "
        "WHERE status!='bekor' AND (wp+mat_sum-paid)>0" + su, pu)
    return {"income": income, "mats": mats, "exp": exp, "done": done, "new": new,
            "debt": debt, "profit": income - mats - exp}


async def month_report_text(month: str, usta=None, admin=False) -> str:
    s = await summary(month, usta)
    t = (
        f"📊 <b>Hisobot: {month}</b>\n\n"
        f"🆕 Yangi zakazlar: {s['new']}\n"
        f"✅ Tugallangan zakazlar: {s['done']}\n\n"
        f"💰 Tushum (to'lovlar): {fmt(s['income'])}\n"
        f"🧱 Materiallar xarajati: {fmt(s['mats'])}\n"
        f"💸 Boshqa xarajatlar: {fmt(s['exp'])}\n"
        f"📈 Sof foyda: <b>{fmt(s['profit'])}</b>\n\n"
        f"💳 Jami qarz (hozirgi): <b>{fmt(s['debt'])}</b>"
    )
    if admin:
        rows = await db.fetchall("SELECT DISTINCT usta_id FROM orders WHERE usta_id IS NOT NULL")
        if rows:
            t += "\n\n👷 <b>Ustalar bo'yicha:</b>"
            for r in rows:
                us = await summary(month, r["usta_id"])
                name = await db.user_name(r["usta_id"])
                t += f"\n• {name}: tugallangan {us['done']}, tushum {fmt(us['income'])}, qarz {fmt(us['debt'])}"
    return t


def _sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[chr(64 + i)].width = max(14, len(h) + 4)


def _flt(col, month, ucol, usta):
    sql, p = "", []
    if month:
        sql += f" AND strftime('%Y-%m', {col})=?"
        p.append(month)
    if usta:
        sql += f" AND {ucol}=?"
        p.append(usta)
    return sql, p


async def build_excel(month=None, usta=None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    s, p = _flt("o.created_at", month, "o.usta_id", usta)
    orders = await db.fetchall(
        f"SELECT o.*, o.wp+o.mat_sum AS total, o.wp+o.mat_sum-o.paid AS debt, "
        f"(SELECT name FROM users u WHERE u.tg_id=o.usta_id) AS usta_name "
        f"FROM ({db.ORDER_SELECT}) o WHERE 1=1{s} ORDER BY o.id", p)
    _sheet(wb, "Zakazlar",
           ["№", "Sana", "Mijoz", "Telefon", "Usta", "Holat", "Tavsif", "Ish haqi", "Materiallar", "Jami", "To'langan", "Qarz"],
           [[o["id"], (o["created_at"] or "")[:16], o["client_name"], o["client_phone"], o["usta_name"] or "",
             STATUS_LABEL.get(o["status"], o["status"]), o["description"], o["wp"], o["mat_sum"],
             o["total"], o["paid"], o["debt"]] for o in orders])

    s, p = _flt("p.created_at", month, "o.usta_id", usta)
    pays = await db.fetchall(
        f"SELECT p.created_at, p.order_id, o.client_name, p.amount, p.kind, p.method FROM payments p "
        f"JOIN orders o ON o.id=p.order_id WHERE 1=1{s} ORDER BY p.id", p)
    _sheet(wb, "To'lovlar", ["Sana", "Zakaz №", "Mijoz", "Summa", "Turi", "Usul"],
           [[x["created_at"][:16], x["order_id"], x["client_name"], x["amount"], x["kind"], x["method"]] for x in pays])

    s, p = _flt("m.created_at", month, "o.usta_id", usta)
    mats = await db.fetchall(
        f"SELECT m.created_at, m.order_id, m.name, m.qty, m.price FROM materials m "
        f"JOIN orders o ON o.id=m.order_id WHERE 1=1{s} ORDER BY m.id", p)
    _sheet(wb, "Materiallar", ["Sana", "Zakaz №", "Nomi", "Miqdori", "Narxi", "Jami"],
           [[x["created_at"][:16], x["order_id"], x["name"], x["qty"], x["price"], x["qty"] * x["price"]] for x in mats])

    s, p = _flt("e.created_at", month, "e.by_user", usta)
    exps = await db.fetchall(
        f"SELECT e.created_at, e.order_id, e.category, e.amount, e.note FROM expenses e WHERE 1=1{s} ORDER BY e.id", p)
    _sheet(wb, "Xarajatlar", ["Sana", "Zakaz №", "Toifa", "Summa", "Izoh"],
           [[x["created_at"][:16], x["order_id"] or "", x["category"], x["amount"], x["note"] or ""] for x in exps])

    su, pu = _u("usta_id", usta)
    debtors = await db.fetchall(
        f"SELECT MAX(client_name) AS name, client_phone, SUM(wp+mat_sum-paid) AS debt "
        f"FROM ({db.ORDER_SELECT}) o WHERE status!='bekor'{su} GROUP BY client_phone "
        f"HAVING debt>0.5 ORDER BY debt DESC", pu)
    _sheet(wb, "Qarzdorlar", ["Mijoz", "Telefon", "Qarz"], [[d["name"], d["client_phone"], d["debt"]] for d in debtors])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
